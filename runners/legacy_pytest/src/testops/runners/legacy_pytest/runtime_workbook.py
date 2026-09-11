"""Materialize the immutable RunSnapshot into the legacy Excel runtime template."""

from __future__ import annotations

import hashlib
import json
import os
from copy import copy
from pathlib import Path
from tempfile import NamedTemporaryFile

from testops.contracts import CaseBaselineRef, CaseDefinition

RUNTIME_WORKSHEET = "自动化测试用例"
REQUIRED_RUNTIME_COLUMNS = frozenset(
    {
        "用例ID",
        "模块",
        "测试场景",
        "测试点",
        "优先级",
        "前置条件",
        "操作步骤",
        "元素定位器",
        "操作类型",
        "输入数据",
        "数据类型",
        "期望结果",
        "验证点",
        "断言类型",
        "超时(秒)",
        "备注",
        "是否执行",
    }
)


class RuntimeWorkbookError(RuntimeError):
    """The package workbook cannot safely represent the requested snapshot."""


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, (list, dict)):
        return json.dumps(value, ensure_ascii=False, separators=(",", ":"), sort_keys=True)
    return str(value)


def _input_data(case: CaseDefinition) -> str:
    return "|".join(_cell_text(step.input) for step in case.steps if step.input is not None)


def _locator_slots(case: CaseDefinition, *, explicit_assertion_column: bool) -> str:
    locators = [step.locator or "" for step in case.steps]
    assertion_locator = case.assertion.locator or ""
    last_locator = next((item for item in reversed(locators) if item), "")
    if assertion_locator and not explicit_assertion_column and assertion_locator != last_locator:
        locators.append(assertion_locator)
    return ", ".join(locators)


def _runtime_values(
    case: CaseDefinition,
    *,
    explicit_assertion_column: bool,
    fallback_execution_group: str,
) -> dict[str, object]:
    values: dict[str, object] = {
        "用例ID": case.case_code,
        "模块": case.module_name,
        "测试场景": case.title,
        "测试点": case.test_point,
        "优先级": case.priority.value,
        "前置条件": "\n".join(case.preconditions),
        "操作步骤": case.source_instructions or "",
        "元素定位器": _locator_slots(
            case,
            explicit_assertion_column=explicit_assertion_column,
        ),
        "操作类型": ",".join(step.operation for step in case.steps),
        "输入数据": _input_data(case),
        "数据类型": case.data_type or "",
        "期望结果": case.expected_result or "",
        "验证点": _cell_text(case.assertion.expected),
        "断言类型": case.assertion.type,
        "超时(秒)": case.timeout_seconds,
        "备注": case.notes or "",
        "实际结果": "",
        "是否执行": "是" if case.enabled else "否",
        "执行分组": case.execution_group or fallback_execution_group or "AUTO",
    }
    if explicit_assertion_column:
        values["断言定位器"] = case.assertion.locator or ""
    return values


def _sha256(path: Path) -> str:
    return f"sha256:{hashlib.sha256(path.read_bytes()).hexdigest()}"


def materialize_runtime_workbook(
    workbook_path: Path,
    cases: tuple[CaseDefinition, ...],
    baseline: CaseBaselineRef,
    *,
    evidence_path: Path,
) -> dict[str, object]:
    """Replace template rows with the exact snapshot cases and record safe evidence."""

    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - the dedicated image pins openpyxl
        raise RuntimeWorkbookError("legacy runtime requires openpyxl") from exc

    try:
        workbook = load_workbook(workbook_path)
    except Exception as exc:
        raise RuntimeWorkbookError("legacy package workbook cannot be opened") from exc

    try:
        worksheet = (
            workbook[RUNTIME_WORKSHEET]
            if RUNTIME_WORKSHEET in workbook.sheetnames
            else workbook.active
        )
        headers = {
            str(worksheet.cell(1, column).value).strip(): column
            for column in range(1, worksheet.max_column + 1)
            if worksheet.cell(1, column).value
        }
        missing = sorted(REQUIRED_RUNTIME_COLUMNS - headers.keys())
        if missing:
            raise RuntimeWorkbookError(
                "legacy package workbook is missing runtime columns: " + ", ".join(missing)
            )

        id_column = headers["用例ID"]
        template_groups = {
            str(worksheet.cell(row, id_column).value or "").strip(): str(
                worksheet.cell(row, headers["执行分组"]).value or ""
            ).strip()
            for row in range(2, worksheet.max_row + 1)
            if "执行分组" in headers and worksheet.cell(row, id_column).value
        }
        template_styles = {
            column: copy(worksheet.cell(2, column)._style)
            for column in range(1, worksheet.max_column + 1)
            if worksheet.max_row >= 2
        }
        template_row_height = worksheet.row_dimensions[2].height if worksheet.max_row >= 2 else None

        if worksheet.max_row > 1:
            worksheet.delete_rows(2, worksheet.max_row - 1)
        explicit_assertion_column = "断言定位器" in headers
        for row_number, case in enumerate(cases, start=2):
            values = _runtime_values(
                case,
                explicit_assertion_column=explicit_assertion_column,
                fallback_execution_group=template_groups.get(case.case_code, ""),
            )
            for header, column in headers.items():
                if header in values:
                    cell = worksheet.cell(row_number, column)
                    cell.value = values[header]
                    if column in template_styles:
                        cell._style = copy(template_styles[column])
            if template_row_height is not None:
                worksheet.row_dimensions[row_number].height = template_row_height

        temporary_path: Path | None = None
        try:
            with NamedTemporaryFile(
                prefix="runtime-test-case-",
                suffix=".xlsx",
                dir=workbook_path.parent,
                delete=False,
            ) as temporary:
                temporary_path = Path(temporary.name)
            workbook.save(temporary_path)
            os.replace(temporary_path, workbook_path)
            temporary_path = None
        finally:
            if temporary_path is not None:
                temporary_path.unlink(missing_ok=True)
    finally:
        workbook.close()

    verification = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        worksheet = (
            verification[RUNTIME_WORKSHEET]
            if RUNTIME_WORKSHEET in verification.sheetnames
            else verification.active
        )
        headers = {
            str(worksheet.cell(1, column).value).strip(): column
            for column in range(1, worksheet.max_column + 1)
            if worksheet.cell(1, column).value
        }
        actual_codes = tuple(
            str(worksheet.cell(row, headers["用例ID"]).value or "").strip()
            for row in range(2, worksheet.max_row + 1)
            if worksheet.cell(row, headers["用例ID"]).value
        )
    finally:
        verification.close()
    expected_codes = tuple(case.case_code for case in cases)
    if actual_codes != expected_codes:
        raise RuntimeWorkbookError("runtime workbook case order does not match RunSnapshot")

    evidence: dict[str, object] = {
        "schema_version": "testops-runtime-workbook-v1",
        "source": "RunSnapshot.cases",
        "baseline_version": baseline.version,
        "baseline_digest": baseline.digest,
        "case_count": len(cases),
        "case_codes": list(expected_codes),
        "workbook_digest": _sha256(workbook_path),
    }
    evidence_path.write_text(
        json.dumps(evidence, ensure_ascii=False, separators=(",", ":"), sort_keys=True),
        encoding="utf-8",
    )
    return evidence
