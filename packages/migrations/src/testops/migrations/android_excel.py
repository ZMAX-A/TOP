"""Import the YanJia Android workbook as managed Excel-case contracts."""

from __future__ import annotations

import hashlib
import math
from collections import Counter
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any
from uuid import UUID, uuid5

from testops.contracts import (
    AssertionDefinition,
    CaseBaseline,
    CaseBaselineSource,
    CaseDefinition,
    CasePriority,
    CaseSourceTrace,
    StepDefinition,
)

from .legacy_excel import write_migration_result

DEFAULT_CASE_WORKSHEET = "自动化测试用例"
DEFAULT_STEP_WORKSHEET = "自动化执行步骤"
ANDROID_BASELINE_NAMESPACE = UUID("42025173-ca28-5bb0-90d1-0ea6ee4d4d2a")
REQUIRED_COLUMNS = (
    "用例ID",
    "模块",
    "测试场景",
    "测试点",
    "优先级",
    "超时(秒)",
    "自动化状态",
    "标签",
    "是否执行",
)
MODULE_KEYS = {
    "账号登录": "login",
    "登录": "login",
    "首页": "home",
    "首页搜索": "home-search",
    "首页跳转": "home-navigation",
    "顾客列表": "customer-list",
    "顾客详情": "customer-detail",
    "影像阅览": "image-viewer",
    "案例库": "case-library",
    "个人中心": "profile",
}
TRUE_VALUES = {"1", "true", "yes", "y", "on", "是", "启用", "执行"}


class AndroidExcelMigrationError(ValueError):
    """The Android workbook cannot be converted without ambiguity."""


@dataclass(frozen=True, slots=True)
class AndroidMigrationResult:
    baseline: CaseBaseline
    audit: dict[str, Any]


def _text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if not math.isfinite(value):
            raise AndroidExcelMigrationError("Excel contains a non-finite numeric value")
        return str(int(value)) if value.is_integer() else format(value, ".15g")
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return str(value).strip()


def _enabled(value: object) -> bool:
    return _text(value).lower() in TRUE_VALUES


def _timeout(value: object) -> float:
    text = _text(value)
    if not text:
        return 10.0
    try:
        timeout = float(text)
    except ValueError as exc:
        raise AndroidExcelMigrationError(f"超时(秒)不是数字：{text}") from exc
    if not 0 < timeout <= 600:
        raise AndroidExcelMigrationError(f"超时(秒)超出范围：{text}")
    return timeout


def _module_key(module_name: str) -> str:
    if module_name in MODULE_KEYS:
        return MODULE_KEYS[module_name]
    digest = hashlib.sha256(module_name.encode("utf-8")).hexdigest()[:10]
    return f"android-{digest}"


def _headers(values: tuple[object, ...]) -> list[str]:
    headers = [_text(value) or f"__blank_column_{index}" for index, value in enumerate(values, 1)]
    missing = [name for name in REQUIRED_COLUMNS if name not in headers]
    if missing:
        raise AndroidExcelMigrationError("Excel 缺少必需列：" + ", ".join(missing))
    return headers


def _tags(raw: object) -> tuple[str, ...]:
    tags = [item.strip().lower() for item in _text(raw).split(",") if item.strip()]
    tags.extend(("android", "excel-driven"))
    if not set(tags).intersection({"mutating", "destructive", "persistent"}):
        tags.append("readonly")
    return tuple(dict.fromkeys(tags))


def _priority(raw: object) -> CasePriority:
    value = _text(raw).upper()
    try:
        return CasePriority(value)
    except ValueError as exc:
        raise AndroidExcelMigrationError(f"不支持的优先级：{value or '<blank>'}") from exc


def _execution_group(raw: object) -> str | None:
    value = _text(raw).upper()
    aliases = {"自动": "AUTO", "串行": "SERIAL", "WORKER_A": "A", "WORKER_B": "B"}
    normalized = aliases.get(value, value)
    if not normalized:
        return None
    if normalized not in {"A", "B", "SERIAL", "AUTO"}:
        raise AndroidExcelMigrationError(f"不支持的执行分组：{value}")
    return normalized


def _counter(values: list[str]) -> dict[str, int]:
    return dict(sorted(Counter(values).items()))


def migrate_android_excel(
    source: str | Path,
    *,
    project_key: str,
    version: str,
    worksheet: str = DEFAULT_CASE_WORKSHEET,
    step_worksheet: str = DEFAULT_STEP_WORKSHEET,
) -> AndroidMigrationResult:
    """Read the source workbook without modifying or copying it."""

    source_path = Path(source)
    if not source_path.is_file():
        raise AndroidExcelMigrationError(f"源 Excel 不存在：{source_path}")
    source_bytes = source_path.read_bytes()
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover
        raise AndroidExcelMigrationError("缺少迁移依赖 openpyxl") from exc

    workbook = load_workbook(source_path, read_only=True, data_only=True)
    try:
        if worksheet not in workbook.sheetnames:
            raise AndroidExcelMigrationError(f"Excel 中不存在工作表「{worksheet}」")
        sheet = workbook[worksheet]
        rows = sheet.iter_rows(values_only=True)
        try:
            headers = _headers(next(rows))
        except StopIteration as exc:
            raise AndroidExcelMigrationError("Excel 主表为空") from exc

        step_counts: Counter[str] = Counter()
        if step_worksheet in workbook.sheetnames:
            step_rows = workbook[step_worksheet].iter_rows(values_only=True)
            try:
                step_headers = [_text(value) for value in next(step_rows)]
            except StopIteration:
                step_headers = []
            if "用例ID" in step_headers:
                case_index = step_headers.index("用例ID")
                enabled_index = (
                    step_headers.index("是否执行") if "是否执行" in step_headers else None
                )
                for values in step_rows:
                    case_code = _text(values[case_index] if case_index < len(values) else None)
                    enabled = (
                        True
                        if enabled_index is None
                        else _enabled(
                            values[enabled_index] if enabled_index < len(values) else None
                        )
                    )
                    if case_code and enabled:
                        step_counts[case_code] += 1

        cases: list[CaseDefinition] = []
        seen: set[str] = set()
        statuses: list[str] = []
        source_row_count = 0
        for row_number, values in enumerate(rows, start=2):
            row = dict(zip(headers, values, strict=False))
            case_code = _text(row.get("用例ID")).upper()
            if not case_code:
                continue
            source_row_count += 1
            if case_code in seen:
                raise AndroidExcelMigrationError(f"用例ID重复：{case_code}")
            seen.add(case_code)
            module_name = _text(row.get("模块")) or "Android"
            automation_status = _text(row.get("自动化状态"))
            statuses.append(automation_status or "<blank>")
            # Match the APP repository: the explicit execution switch is authoritative.
            # Some migrated Android cases retain a historical BACKLOG label after their
            # executable steps have been added, so status text must not silently drop them.
            enabled = _enabled(row.get("是否执行"))
            tags = _tags(row.get("标签"))
            precondition = _text(row.get("前置条件"))
            expected_result = _text(row.get("期望结果"))
            cases.append(
                CaseDefinition(
                    case_id=uuid5(
                        ANDROID_BASELINE_NAMESPACE,
                        f"case:{project_key}:{case_code}",
                    ),
                    case_code=case_code,
                    module_key=_module_key(module_name),
                    module_name=module_name,
                    title=_text(row.get("测试场景")) or case_code,
                    test_point=_text(row.get("测试点")) or case_code,
                    priority=_priority(row.get("优先级")),
                    preconditions=(precondition,) if precondition else (),
                    steps=(StepDefinition(operation="execute_excel_case", input=case_code),),
                    assertion=AssertionDefinition(type="excel_case_passed", expected="PASS"),
                    tags=tags,
                    enabled=enabled,
                    source_instructions=_text(row.get("操作步骤")) or None,
                    data_type=_text(row.get("数据类型")) or None,
                    expected_result=expected_result or None,
                    timeout_seconds=_timeout(row.get("超时(秒)")),
                    source_trace=CaseSourceTrace(worksheet=worksheet, row_number=row_number),
                    notes=_text(row.get("备注")) or None,
                    execution_group=_execution_group(row.get("执行分组")),
                )
            )
    finally:
        workbook.close()

    if not cases:
        raise AndroidExcelMigrationError("Excel 中没有用例")
    source_digest = f"sha256:{hashlib.sha256(source_bytes).hexdigest()}"
    baseline = CaseBaseline(
        baseline_id=uuid5(
            ANDROID_BASELINE_NAMESPACE,
            f"baseline:{project_key}:{version}",
        ),
        project_key=project_key,
        version=version,
        source=CaseBaselineSource(
            file_name=source_path.name,
            file_digest=source_digest,
            worksheet=worksheet,
        ),
        cases=tuple(cases),
    )
    audit: dict[str, Any] = {
        "schema_version": "1.0",
        "migration": "yanjia-android-excel-to-managed-case-baseline",
        "source": {
            "file_name": source_path.name,
            "file_digest": source_digest,
            "worksheet": worksheet,
            "step_worksheet": step_worksheet,
        },
        "target": {
            "baseline_id": str(baseline.baseline_id),
            "project_key": project_key,
            "version": version,
        },
        "policies": {
            "execution": "one managed execute_excel_case operation per case code",
            "secrets": "workbook copied neither to the baseline nor to the control plane",
            "immutability": "existing baseline files may only be reused byte-for-byte",
        },
        "counts": {
            "source_cases": source_row_count,
            "baseline_cases": len(cases),
            "enabled_cases": sum(case.enabled for case in cases),
            "disabled_cases": sum(not case.enabled for case in cases),
            "enabled_step_rows": sum(step_counts.values()),
            "cases_with_step_rows": len(step_counts),
        },
        "distribution": {
            "automation_status": _counter(statuses),
            "modules": _counter([case.module_name for case in cases]),
            "priorities": _counter([case.priority.value for case in cases]),
            "safety_tags": _counter(
                [
                    tag
                    for case in cases
                    for tag in case.tags
                    if tag in {"readonly", "mutating", "destructive", "persistent", "no_retry"}
                ]
            ),
        },
        "step_rows_by_case": dict(sorted(step_counts.items())),
        "errors": [],
    }
    return AndroidMigrationResult(baseline=baseline, audit=audit)


__all__ = [
    "AndroidExcelMigrationError",
    "AndroidMigrationResult",
    "DEFAULT_CASE_WORKSHEET",
    "DEFAULT_STEP_WORKSHEET",
    "migrate_android_excel",
    "write_migration_result",
]
