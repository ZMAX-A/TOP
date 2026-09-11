from __future__ import annotations

import json
import shutil
import tempfile
from pathlib import Path
from uuid import uuid4

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill

from testops.contracts import (
    AssertionDefinition,
    CaseBaselineRef,
    CaseDefinition,
    CasePriority,
    StepDefinition,
)
from testops.migrations.legacy_excel import migrate_legacy_excel
from testops.runners.legacy_pytest.runtime_workbook import materialize_runtime_workbook

ROOT = Path(__file__).resolve().parents[2]
CURRENT_WEB_WORKBOOK = ROOT.parent / "web_test" / "web" / "test_cases" / "test_case.xlsx"

HEADERS = (
    "用例ID",
    "模块",
    "测试场景",
    "测试点",
    "优先级",
    "前置条件",
    "操作步骤",
    "元素定位器",
    "操作类型",
    "输入数据",
    "数据类型",
    "期望结果",
    "验证点",
    "断言类型",
    "超时(秒)",
    "备注",
    "实际结果",
    "是否执行",
    "执行分组",
)


def _template(path: Path) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "自动化测试用例"
    worksheet.append(HEADERS)
    worksheet.append(
        (
            "TC-DETAIL-001",
            "旧模块",
            "旧场景",
            "旧测试点",
            "P3",
            "",
            "",
            "#stale",
            "click",
            "",
            "",
            "旧预期",
            "旧验证点",
            "text_visible",
            5,
            "",
            "旧结果",
            "是",
            "B",
        )
    )
    worksheet["A2"].fill = PatternFill(fill_type="solid", fgColor="00DDEEFF")
    workbook.save(path)
    workbook.close()


def _case(*, execution_group: str | None = None) -> CaseDefinition:
    return CaseDefinition(
        case_id=uuid4(),
        case_code="TC-DETAIL-001",
        module_key="customer-detail",
        module_name="顾客详情",
        title="新场景",
        test_point="新测试点",
        priority=CasePriority.P1,
        preconditions=("已登录",),
        source_instructions="1. 输入筛选词\n2. 点击结果",
        steps=(
            StepDefinition(operation="input", locator="#search", input="测试"),
            StepDefinition(operation="click", locator="button[type=submit]"),
        ),
        assertion=AssertionDefinition(
            type="text_equals",
            expected="新预期",
            locator="#result",
        ),
        data_type="文本",
        expected_result="展示新预期",
        timeout_seconds=20,
        notes="运行时快照",
        execution_group=execution_group,
    )


def _baseline() -> CaseBaselineRef:
    return CaseBaselineRef(
        baseline_id=uuid4(),
        version="case-v1.0.6",
        digest="sha256:" + "a" * 64,
        case_count=1,
    )


def test_materializes_snapshot_values_and_preserves_template_execution_group() -> None:
    with tempfile.TemporaryDirectory() as temporary:
        root = Path(temporary)
        workbook_path = root / "test_case.xlsx"
        evidence_path = root / "runtime-workbook.json"
        _template(workbook_path)

        evidence = materialize_runtime_workbook(
            workbook_path,
            (_case(),),
            _baseline(),
            evidence_path=evidence_path,
        )

        workbook = load_workbook(workbook_path, read_only=False, data_only=True)
        worksheet = workbook["自动化测试用例"]
        headers = {cell.value: cell.column for cell in worksheet[1]}
        row = {name: worksheet.cell(2, column).value for name, column in headers.items()}
        fill_type = worksheet["A2"].fill.fill_type
        fill_color = worksheet["A2"].fill.fgColor.rgb
        workbook.close()
        evidence_document = json.loads(evidence_path.read_text("utf-8"))

    assert row["模块"] == "顾客详情"
    assert row["测试场景"] == "新场景"
    assert row["元素定位器"] == "#search, button[type=submit], #result"
    assert row["操作类型"] == "input,click"
    assert row["输入数据"] == "测试"
    assert row["验证点"] == "新预期"
    assert row["实际结果"] is None
    assert row["执行分组"] == "B"
    assert fill_type == "solid"
    assert fill_color == "00DDEEFF"
    assert evidence["source"] == "RunSnapshot.cases"
    assert evidence["case_codes"] == ["TC-DETAIL-001"]
    assert evidence_document == evidence


def test_snapshot_execution_group_overrides_package_template() -> None:
    with tempfile.TemporaryDirectory() as temporary:
        root = Path(temporary)
        workbook_path = root / "test_case.xlsx"
        _template(workbook_path)

        materialize_runtime_workbook(
            workbook_path,
            (_case(execution_group="SERIAL"),),
            _baseline(),
            evidence_path=root / "runtime-workbook.json",
        )

        workbook = load_workbook(workbook_path, read_only=True, data_only=True)
        worksheet = workbook["自动化测试用例"]
        headers = {cell.value: cell.column for cell in worksheet[1]}
        execution_group = worksheet.cell(2, headers["执行分组"]).value
        workbook.close()

    assert execution_group == "SERIAL"


@pytest.mark.skipif(not CURRENT_WEB_WORKBOOK.exists(), reason="Web workbook is outside this repo")
def test_current_web_workbook_round_trips_through_runtime_snapshot() -> None:
    source = migrate_legacy_excel(
        CURRENT_WEB_WORKBOOK,
        project_key="yanjia-ai-web",
        version="case-v1.0.6",
    ).baseline
    baseline_ref = CaseBaselineRef(
        baseline_id=source.baseline_id,
        version=source.version,
        digest="sha256:" + "b" * 64,
        case_count=len(source.cases),
    )

    with tempfile.TemporaryDirectory() as temporary:
        root = Path(temporary)
        runtime_workbook = root / "test_case.xlsx"
        shutil.copy2(CURRENT_WEB_WORKBOOK, runtime_workbook)
        materialize_runtime_workbook(
            runtime_workbook,
            source.cases,
            baseline_ref,
            evidence_path=root / "runtime-workbook.json",
        )
        round_tripped = migrate_legacy_excel(
            runtime_workbook,
            project_key="yanjia-ai-web",
            version="case-v1.0.6",
        ).baseline

    def semantic_cases(cases):
        return [
            case.model_dump(mode="json", exclude={"source_trace"}, exclude_none=True)
            for case in cases
        ]

    assert semantic_cases(round_tripped.cases) == semantic_cases(source.cases)
